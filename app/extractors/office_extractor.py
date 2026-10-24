from docx import Document
from pptx import Presentation
from pathlib import Path
from app.schemas.extraction import PageResult
from app.utils.normalize import normalize_text
from app.utils.logger import logger


def extract_docx(file_path: str) -> list:
    logger.info(f"Procesando DOCX: {file_path}")
    warnings = []
    try:
        doc = Document(file_path)
        texts = []
        for para in doc.paragraphs:
            if para.text.strip():
                texts.append(para.text)
        text_final = normalize_text('\n\n'.join(texts))
        if not text_final:
            warnings.append("No se encontró texto en el documento")
        return [PageResult(
            page_number=1,
            text=text_final,
            method="text",
            has_image=False,
            warnings=warnings
        )]
    except Exception as e:
        return [PageResult(
            page_number=1,
            text="",
            method="text",
            has_image=False,
            warnings=[f"Error al procesar DOCX: {str(e)}"]
        )]


def extract_pptx(file_path: str) -> list:
    logger.info(f"Procesando PPTX: {file_path}")
    warnings = []
    try:
        prs = Presentation(file_path)
        pages_results = []
        for idx, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text)
            text_final = normalize_text('\n\n'.join(texts))
            pages_results.append(PageResult(
                page_number=idx + 1,
                text=text_final,
                method="text",
                has_image=True,
                warnings=[]
            ))
        if not any(p.text for p in pages_results):
            warnings.append("No se encontró texto en las diapositivas")
        # Añadir warning general si aplica
        if warnings and pages_results:
            pages_results[0].warnings = warnings
        return pages_results
    except Exception as e:
        return [PageResult(
            page_number=1,
            text="",
            method="text",
            has_image=True,
            warnings=[f"Error al procesar PPTX: {str(e)}"]
        )]


def extract_office(file_path: str) -> list:
    ext = Path(file_path).suffix.lower()
    if ext == '.docx':
        return extract_docx(file_path)
    elif ext == '.pptx':
        return extract_pptx(file_path)
    return []
